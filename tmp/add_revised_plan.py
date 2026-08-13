"""Add a 'Revised Plan' sheet to the existing Gantt workbook.

Reflects status snapshot as of 2026-05-21 (Week 10 of the original plan).
Total remaining budget: 408 hrs across Ayu + Asmerom (Tushar excluded — FCB resource).
"""
from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

WB_PATH = "docs/Flink_K8s_CP_Gantt_Chart.xlsx"
SHEET_NAME = "Revised Plan — May 21"

NAVY = "FF0A2F5C"
NAVY_DARK = "FF003762"
WHITE = "FFFFFFFF"
LIGHT_GREY = "FFF2F2F2"
DONE = "FFC8E6C9"        # green
IN_PROGRESS = "FFFFE0B2"  # amber
BLOCKED = "FFFFCDD2"      # red
PENDING = "FFE0E0E0"      # grey
FLINK_BAND = "FFE3F2FD"   # light blue
CFK_BAND = "FFF3E5F5"     # light purple

thin = Side(border_style="thin", color="FFBDBDBD")
border = Border(left=thin, right=thin, top=thin, bottom=thin)


def hdr_font():
    return Font(name="Calibri", size=11, bold=True, color=WHITE)


def section_font():
    return Font(name="Calibri", size=11, bold=True, color=WHITE)


def title_font():
    return Font(name="Calibri", size=14, bold=True, color=NAVY_DARK)


def body_font(bold=False):
    return Font(name="Calibri", size=11, bold=bold)


def fill(color):
    return PatternFill(fill_type="solid", fgColor=color)


def write_row(ws, row, values, fills=None, fonts=None, aligns=None):
    for i, v in enumerate(values, start=1):
        c = ws.cell(row=row, column=i, value=v)
        if fills and i - 1 < len(fills) and fills[i - 1]:
            c.fill = fills[i - 1]
        if fonts and i - 1 < len(fonts) and fonts[i - 1]:
            c.font = fonts[i - 1]
        if aligns and i - 1 < len(aligns) and aligns[i - 1]:
            c.alignment = aligns[i - 1]
        c.border = border


def main():
    wb = load_workbook(WB_PATH)
    if SHEET_NAME in wb.sheetnames:
        del wb[SHEET_NAME]
    ws = wb.create_sheet(SHEET_NAME)

    # --- column widths
    widths = {"A": 5, "B": 48, "C": 14, "D": 14, "E": 12, "F": 60}
    for col, w in widths.items():
        ws.column_dimensions[col].width = w

    center = Alignment(horizontal="center", vertical="center", wrap_text=True)
    left = Alignment(horizontal="left", vertical="center", wrap_text=True)

    # --- title
    ws.merge_cells("A1:F1")
    t = ws["A1"]
    t.value = "Revised Plan — Snapshot 2026-05-21 (Week 10)"
    t.font = title_font()
    t.alignment = Alignment(horizontal="left", vertical="center")
    ws.row_dimensions[1].height = 24

    # --- budget summary
    ws.merge_cells("A2:F2")
    s = ws["A2"]
    s.value = (
        "Remaining budget: 408 hrs across Ayu + Asmerom. "
        "Tushar = FCB resource (no allocation against this pool). "
        "Hour estimates are tentative — adjust to fit actual read."
    )
    s.font = Font(name="Calibri", size=10, italic=True, color="FF555555")
    s.alignment = left
    ws.row_dimensions[2].height = 30

    # --- column header row
    headers = ["#", "Activity", "Owner", "Status", "Est. Hrs", "Notes"]
    write_row(
        ws,
        4,
        headers,
        fills=[fill(NAVY_DARK)] * 6,
        fonts=[hdr_font()] * 6,
        aligns=[center] * 6,
    )
    ws.row_dimensions[4].height = 22

    status_fill = {
        "Done": fill(DONE),
        "In Progress": fill(IN_PROGRESS),
        "Blocked": fill(BLOCKED),
        "Pending": fill(PENDING),
    }

    # --- data: track sections
    # (activity, owner, status, hrs, notes)
    flink_rows = [
        (
            "All-env Flink manifests (DEV/QA/UAT/PRD)",
            "Ayu",
            "Done",
            0,
            "Manifests authored for every env; UAT/PRD apply still pending cluster provision",
        ),
        (
            "Flink Infra CI + Job CI/CD pipelines",
            "Ayu",
            "Done",
            0,
            "One-time build; reused across all envs",
        ),
        (
            "Provision Flink DEV cluster",
            "Ayu",
            "Done",
            0,
            "Live; CMF/Env/Catalog points at VM CP per migration plan",
        ),
        (
            "Provision Flink QA cluster",
            "Tushar",
            "Done",
            0,
            "FCB resource — no allocation",
        ),
        (
            "Provision Flink UAT cluster",
            "Ayu",
            "Pending",
            16,
            "Apply existing manifest + smoke test",
        ),
        (
            "Provision Flink PRD cluster",
            "Ayu",
            "Pending",
            20,
            "Apply existing manifest + smoke test + change-mgmt window",
        ),
        (
            "Flink Job #1 — develop & deploy",
            "Asmerom",
            "Done",
            0,
            "Built and deployed",
        ),
        (
            "Flink Job #2 — develop & deploy",
            "Asmerom",
            "In Progress",
            96,
            "In active development; includes integration test + deploy across envs",
        ),
        (
            "Flink Job #2 — hardening & buffer",
            "Asmerom",
            "Pending",
            32,
            "Defect fix, performance pass, contingency",
        ),
    ]

    cfk_rows = [
        (
            "CFK DEV manifests + DevOps",
            "Ayu",
            "Done",
            0,
            "Kafka / KRaft / KRP / SR CRs authored; CI pipelines in place",
        ),
        (
            "CFK CI/CD pipelines (one-time, all envs)",
            "Ayu",
            "Done",
            0,
            "Reused across DEV/QA/UAT/STG/PRD1/PRD2",
        ),
        (
            "Deploy CFK DEV cluster",
            "Ayu",
            "Blocked",
            24,
            "Blocked on storage (ResourceQuota / PVC provisioning); unblock then validate",
        ),
        (
            "CFK STG manifests + DevOps",
            "Ayu",
            "Pending",
            20,
            "Clone DEV manifests, apply STG overrides",
        ),
        (
            "CFK UAT manifests + DevOps",
            "Ayu",
            "Pending",
            20,
            "Clone DEV manifests, apply UAT overrides",
        ),
        (
            "CFK PRD1 manifests + DevOps",
            "Ayu",
            "Pending",
            24,
            "Production sizing + HA review",
        ),
        (
            "CFK DR manifests + DevOps",
            "Ayu",
            "Pending",
            24,
            "Disaster-recovery topology, cluster-link to PRD1",
        ),
        (
            "Deploy CFK QA",
            "Tushar",
            "Pending",
            8,
            "Tushar lead; Ayu consulting hours only",
        ),
        (
            "Deploy CFK UAT",
            "Tushar",
            "Pending",
            8,
            "Tushar lead; Ayu consulting hours only",
        ),
        (
            "Deploy CFK STG",
            "Tushar",
            "Pending",
            8,
            "Tushar lead; Ayu consulting hours only",
        ),
        (
            "Deploy CFK PRD1",
            "Tushar",
            "Pending",
            12,
            "Tushar lead; Ayu consulting hours only",
        ),
        (
            "Deploy CFK PRD2",
            "Tushar",
            "Pending",
            12,
            "Tushar lead; Ayu consulting hours only",
        ),
        (
            "Cluster link + migration — DEV",
            "Ayu",
            "Pending",
            20,
            "VM CP → CFK; mirror topics, cut over consumers",
        ),
        (
            "Cluster link + migration — UAT",
            "Ayu",
            "Pending",
            24,
            "Includes app-team coordination",
        ),
        (
            "Cluster link + migration — STG",
            "Ayu",
            "Pending",
            24,
            "Pre-prod validation gate",
        ),
        (
            "Cluster link + migration — PRD",
            "Ayu",
            "Pending",
            36,
            "Change window, rollback plan, sign-off",
        ),
    ]

    governance_rows = [
        (
            "Documentation, runbooks, handoff",
            "Ayu",
            "Pending",
            16,
            "Operational runbooks + ROSA-CR-OVERRIDES + cluster-link playbook",
        ),
        (
            "Contingency / unscheduled support",
            "Ayu + Asmerom",
            "Pending",
            12,
            "Buffer for SNOW tickets, defects, Rick/Tom review cycles",
        ),
    ]

    def write_section(start_row, label, color, rows):
        # band header
        ws.merge_cells(start_row=start_row, start_column=1,
                       end_row=start_row, end_column=6)
        c = ws.cell(row=start_row, column=1, value=label)
        c.fill = fill(NAVY)
        c.font = section_font()
        c.alignment = Alignment(horizontal="left", vertical="center")
        ws.row_dimensions[start_row].height = 20

        r = start_row + 1
        for i, (act, owner, status, hrs, notes) in enumerate(rows, start=1):
            sfill = status_fill.get(status, fill(WHITE))
            band = fill(color)
            write_row(
                ws,
                r,
                [i, act, owner, status, hrs, notes],
                fills=[band, band, band, sfill, band, band],
                fonts=[body_font(), body_font(), body_font(), body_font(bold=True),
                       body_font(), body_font()],
                aligns=[center, left, center, center, center, left],
            )
            r += 1
        return r

    row = 5
    row = write_section(row, "FLINK TRACK", FLINK_BAND, flink_rows)
    row = write_section(row, "CFK TRACK", CFK_BAND, cfk_rows)
    row = write_section(row, "GOVERNANCE & CONTINGENCY", FLINK_BAND, governance_rows)

    # --- totals
    all_rows = flink_rows + cfk_rows + governance_rows
    by_owner = {"Ayu": 0, "Asmerom": 0, "Tushar": 0, "Ayu + Asmerom": 0}
    for _, owner, _, hrs, _ in all_rows:
        by_owner[owner] = by_owner.get(owner, 0) + hrs

    ayu_total = by_owner.get("Ayu", 0)
    asm_total = by_owner.get("Asmerom", 0)
    pool_total = ayu_total + asm_total + by_owner.get("Ayu + Asmerom", 0)
    tushar_total = by_owner.get("Tushar", 0)

    totals_start = row + 1
    # spacer
    row = totals_start

    def write_total(r, label, val, note=""):
        ws.cell(row=r, column=2, value=label).font = body_font(bold=True)
        ws.cell(row=r, column=2).alignment = left
        ws.cell(row=r, column=5, value=val).font = body_font(bold=True)
        ws.cell(row=r, column=5).alignment = center
        if note:
            ws.cell(row=r, column=6, value=note).font = Font(
                name="Calibri", size=10, italic=True, color="FF555555"
            )
            ws.cell(row=r, column=6).alignment = left

    write_total(row, "Ayu — allocated hrs", ayu_total)
    row += 1
    write_total(row, "Asmerom — allocated hrs", asm_total)
    row += 1
    write_total(row, "Shared buffer — Ayu + Asmerom", by_owner.get("Ayu + Asmerom", 0))
    row += 1
    write_total(
        row,
        "Pool total (must be ≤ 408)",
        pool_total,
        f"Budget 408; remaining {408 - pool_total}",
    )
    row += 1
    write_total(
        row,
        "Tushar (FCB) — informational only",
        tushar_total,
        "Excluded from 408-hr pool",
    )

    # --- legend
    row += 2
    ws.cell(row=row, column=2, value="Legend:").font = body_font(bold=True)
    row += 1
    for label, color in [
        ("Done", DONE),
        ("In Progress", IN_PROGRESS),
        ("Blocked", BLOCKED),
        ("Pending", PENDING),
    ]:
        c = ws.cell(row=row, column=2, value=f"  ■  {label}")
        c.font = body_font()
        c.fill = fill(color)
        row += 1

    # freeze the column header row
    ws.freeze_panes = "A5"

    wb.save(WB_PATH)
    print(
        f"Saved. Ayu={ayu_total} Asm={asm_total} shared={by_owner.get('Ayu + Asmerom', 0)} "
        f"pool={pool_total}/408 Tushar={tushar_total}"
    )


if __name__ == "__main__":
    main()
